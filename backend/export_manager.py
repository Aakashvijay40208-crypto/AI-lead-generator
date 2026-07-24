import io
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

def generate_leads_excel(leads):
    """
    Generate an Excel sheet of B2B leads.
    Returns a BytesIO stream containing the workbook binary.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "OXIS Qualified Leads"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Theme colors
    navy_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # #1E1B4B - very dark indigo
    white_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    
    # Priority Fills
    hot_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light red/rose #FEE2E2
    medium_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light amber #FEF3C7
    low_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid") # Light gray #F3F4F6
    
    # Center & Left alignments
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Borders
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    headers = [
        "Business Name", 
        "Category", 
        "Address", 
        "Phone", 
        "WhatsApp", 
        "Email", 
        "LinkedIn", 
        "Website", 
        "Lead Score", 
        "Priority",
        "Problems Detected", 
        "Recommended Services", 
        "Google Maps URL"
    ]
    
    # Write headers
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = white_font
        cell.alignment = align_center
        cell.border = thin_border
        
    # Write data
    for idx, lead in enumerate(leads, 2):
        ws.row_dimensions[idx].height = 24
        
        audit = lead.get("audit_reports") or {}
        scores = lead.get("lead_scores") or {}
        contacts = lead.get("contacts") or {}
        
        # Phone / WhatsApp
        phone = lead.get("phone_number") or ""
        tech = audit.get("tech_detected") or {}
        whatsapp = phone if tech.get("whatsapp") else ""
        
        # Emails
        emails = contacts.get("emails", [])
        email_str = ", ".join(emails) if emails else ""
        
        # Socials
        socials = contacts.get("socialLinks") or {}
        linkedin = socials.get("linkedin", "")
        
        # Score and priority
        score = scores.get("score", 0)
        priority = scores.get("priority", "LOW")
        
        # Problems and Services
        problems = scores.get("failed_checks", [])
        problems_str = ", ".join(problems) if problems else ""
        
        services = scores.get("recommended_services", [])
        services_str = ", ".join(services) if services else ""
        
        # Construct google maps url
        gmaps_url = ""
        place_id = lead.get("place_id")
        lat = lead.get("latitude")
        lng = lead.get("longitude")
        if lat and lng:
            gmaps_url = f"https://www.google.com/maps/search/?api=1&query={lat},{lng}"
        elif place_id:
            gmaps_url = f"https://www.google.com/maps/place/?q=place_id:{place_id}"
            
        row_data = [
            lead.get("name", "N/A"),
            lead.get("category", "N/A") if lead.get("category") else (problems[0].replace("No ", "") if problems else "General"),
            lead.get("address", "N/A"),
            phone,
            whatsapp,
            email_str,
            linkedin,
            lead.get("website", ""),
            score,
            priority,
            problems_str,
            services_str,
            gmaps_url
        ]
        
        ws.append(row_data)
        
        # Style current row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            
            # Alignments
            if col_idx in [4, 5, 9, 10]:  # Phone, WhatsApp, Score, Priority
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
            # Score font bold
            if col_idx == 9:
                cell.font = bold_font
                
            # Apply color highlights to Priority column and Score column
            if col_idx in [9, 10]:
                if priority == "HOT":
                    cell.fill = hot_fill
                elif priority == "MEDIUM":
                    cell.fill = medium_fill
                else:
                    cell.fill = low_fill
                    
    # Auto-adjust column widths with padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        
        for cell in col:
            # Avoid long formulas or URLs inflating width excessively
            val_str = str(cell.value or "")
            if val_str.startswith("http"):
                val_str = "Link"
            max_len = max(max_len, len(val_str))
            
        # Add buffer
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
        
    # Write to memory stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output
